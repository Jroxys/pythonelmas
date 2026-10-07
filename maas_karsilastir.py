#!/usr/bin/env python3
"""İki Excel dosyasındaki maaşları isme göre karşılaştırır.

Dosya 1: Sayfalara bölünmüş (Anaokulu, İlkokul, Ortaokul, Lise ...) - tüm sayfalar birleştirilir.
Dosya 2: Tek liste (tek sayfa).

Kullanım:
    python maas_karsilastir.py                       # pencere / soru-cevap ile sorar
    python maas_karsilastir.py a.xlsx b.xlsx         # doğrudan çalıştırır
    python maas_karsilastir.py a.xlsx b.xlsx -o sonuc.xlsx --tolerans 1

Gereksinim: pip install openpyxl
"""
import argparse
import difflib
import os
import re
import sys
import unicodedata
from collections import defaultdict

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl kurulu değil. Şunu çalıştırın:  pip install openpyxl")

AD_ANAHTAR = ["adı soyadı", "ad soyad", "adi soyadi", "isim", "personel", "ad", "adı", "çalışan"]
MAAS_ANAHTAR = ["maaş", "maas", "ücret", "ucret", "net", "brüt", "brut", "tutar"]


# ---------------------------------------------------------------- yardımcılar
def tr_kucult(s):
    """Türkçe uyumlu küçük harf (İ->i, I->ı)."""
    return s.replace("İ", "i").replace("I", "ı").lower()


def isim_anahtar(s):
    """Karşılaştırma anahtarı: büyük/küçük harf, fazla boşluk ve Türkçe karakter farkını yok sayar."""
    s = tr_kucult(str(s)).strip()
    s = s.translate(str.maketrans("çğıöşü", "cgiosu"))
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^\w\s]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def isim_uyumlu(a, b):
    """Kısa olan ismin her kelimesi diğer isimde (küçük yazım farkıyla) geçiyorsa True.
    Eksik orta isim ve kelime sırası farkı sorun sayılmaz; 'Ahmet Turan Tosun' / 'Ahsen Tosun' uyumsuzdur."""
    x, y = isim_anahtar(a).split(), isim_anahtar(b).split()
    if len(x) > len(y):
        x, y = y, x
    return all(any(difflib.SequenceMatcher(None, k, t).ratio() >= 0.8 for t in y) for k in x)


ASGARI_MAAS = 28075.0  # maaş hücresine "asgari" yazılırsa bu tutar sayılır (ayarlanabilir)


def asgari_mi(v):
    return isinstance(v, str) and "asgari" in tr_kucult(v).replace("ı", "i")


def sayiya_cevir(v):
    """Hücre değerini sayıya çevirir. '1.234,56' ve '1,234.56' gibi metinleri de anlar.
    Hücrede "asgari" yazıyorsa ASGARI_MAAS değerini döndürür."""
    if v is None or isinstance(v, bool):
        return None
    if asgari_mi(v):
        return ASGARI_MAAS
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.\-]", "", str(v))
    if not s or s in "-.,":
        return None
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif s.count(".") > 1 or re.fullmatch(r"-?\d{1,3}\.\d{3}", s):
        s = s.replace(".", "")  # 1.234 veya 1.234.567 -> binlik ayracı
    try:
        return float(s)
    except ValueError:
        return None


def baslik_bul(satirlar):
    """İlk 20 satırda başlık satırını ve ad / maaş sütun adaylarını bulur."""
    en_iyi = None
    for i, satir in enumerate(satirlar[:20]):
        ad_adaylari, maas_adaylari = [], []
        for j, h in enumerate(satir):
            if h is None or isinstance(h, (int, float)):
                continue
            t = tr_kucult(str(h)).strip()
            if not t:
                continue
            if any(t == k or (len(k) > 3 and k in t) for k in AD_ANAHTAR):
                ad_adaylari.append(j)
            if any(k in t for k in MAAS_ANAHTAR):
                maas_adaylari.append(j)
        if ad_adaylari and maas_adaylari:
            return i, ad_adaylari, maas_adaylari
        if en_iyi is None and (ad_adaylari or maas_adaylari):
            en_iyi = (i, ad_adaylari, maas_adaylari)
    return en_iyi


SECICI = None  # arayüz (GUI) kendi seçim penceresini buraya bağlar: fonksiyon(mesaj, secenekler) -> değer


def sec(mesaj, secenekler):
    """Birden fazla aday varsa kullanıcıya sorar (etkileşimsizse ilkini seçer)."""
    if len(secenekler) == 1:
        return secenekler[0][0]
    if SECICI:
        return SECICI(mesaj, secenekler)
    if sys.stdin is None or not sys.stdin.isatty():
        return secenekler[0][0]
    print(mesaj)
    for n, (_, etiket) in enumerate(secenekler, 1):
        print(f"  {n}) {etiket}")
    while True:
        c = input("Seçiminiz [1]: ").strip() or "1"
        if c.isdigit() and 1 <= int(c) <= len(secenekler):
            return secenekler[int(c) - 1][0]


def sutun_sec(mesaj, tur, adaylar, baslik, secilen):
    """Daha önce seçilen başlık bu sayfada da varsa onu kullanır, yoksa sorar."""
    if tur in secilen:
        for j in adaylar:
            if tr_kucult(str(baslik[j])).strip() == secilen[tur]:
                return j
    j = sec(mesaj, [(j, f"{get_column_letter(j + 1)}: {baslik[j]}") for j in adaylar])
    secilen[tur] = tr_kucult(str(baslik[j])).strip()
    return j


def sutun_bul(baslik, kosul):
    for j, h in enumerate(baslik):
        if h is not None and not isinstance(h, (int, float)) and kosul(isim_anahtar(h)):
            return j
    return None


def tc_temizle(v):
    """11 haneli TC kimlik numarasını döndürür, geçersizse None."""
    if v is None:
        return None
    s = str(int(v)) if isinstance(v, (int, float)) else re.sub(r"\D", "", str(v))
    return s if len(s) == 11 else None


def harf_to_idx(h):
    n = 0
    for ch in h.upper():
        n = n * 26 + ord(ch) - 64
    return n - 1


def kisileri_oku(yol, ad_sutun=None, maas_sutun=None, etiket="Dosya"):
    """Dosyadaki tüm sayfalardan (isim, maaş, sayfa, satır) kayıtlarını okur."""
    wb = load_workbook(yol, data_only=True, read_only=True)
    kayitlar, uyarilar = [], []
    secilen = {}  # ilk sayfada seçilen başlıklar, diğer sayfalarda aynen kullanılır
    for ws in wb.worksheets:
        satirlar = [list(r) for r in ws.iter_rows(values_only=True)]
        if not any(any(c is not None for c in r) for r in satirlar):
            continue
        if ad_sutun and maas_sutun:
            tc_i = elden_i = okul_i = None
            bas, ad_i, maas_i = 0, harf_to_idx(ad_sutun), harf_to_idx(maas_sutun)
            # başlık varsa atla: maaş sütunu sayı olmayan ilk satırlar
            while bas < len(satirlar) and (len(satirlar[bas]) <= maas_i or sayiya_cevir(satirlar[bas][maas_i]) is None):
                bas += 1
            bas -= 1
        else:
            b = baslik_bul(satirlar)
            if not b or not b[1] or not b[2]:
                uyarilar.append(f"[{etiket}] '{ws.title}' sayfasında ad/maaş başlığı bulunamadı, atlandı.")
                continue
            bas, adlar, maaslar = b
            baslik = satirlar[bas]
            ad_i = sutun_sec(f"[{etiket} / {ws.title}] İsim sütunu hangisi?", "ad", adlar, baslik, secilen)
            maas_i = sutun_sec(f"[{etiket} / {ws.title}] Maaş sütunu hangisi?", "maas", maaslar, baslik, secilen)
            okul_i = sutun_bul(baslik, lambda t: t.startswith("okul"))
            tc_i = sutun_bul(baslik, lambda t: "kimlik" in t or re.search(r"\bt ?c\b", t))
            elden_i = sutun_bul(baslik, lambda t: "elden" in t)
        for r_no, satir in enumerate(satirlar[bas + 1:], start=bas + 2):
            if len(satir) <= max(ad_i, maas_i):
                continue
            ad = satir[ad_i]
            if ad is None or not str(ad).strip():
                continue
            asgari = asgari_mi(satir[maas_i])
            maas = sayiya_cevir(satir[maas_i])
            if maas is None:
                continue  # "TOPLAM" gibi sayısal olmayan satırlar
            if isim_anahtar(ad) in ("toplam", "genel toplam", "ara toplam"):
                continue
            tc = tc_temizle(satir[tc_i]) if tc_i is not None and tc_i < len(satir) else None
            elden = None
            if elden_i is not None and elden_i < len(satir):
                elden = sayiya_cevir(satir[elden_i]) or 0.0
            okul = None
            if okul_i is not None and okul_i < len(satir) and satir[okul_i] is not None:
                okul = str(satir[okul_i]).strip()
            kayitlar.append({"ad": str(ad).strip(), "anahtar": isim_anahtar(ad), "tc": tc, "elden": elden,
                             "okul": okul, "maas": maas, "asgari": asgari, "sayfa": ws.title, "satir": r_no})
    return kayitlar, uyarilar


# ---------------------------------------------------------------- karşılaştırma
def _esle(l1, l2, anahtar_fn, yontem, tolerans, elden_dahil, eslesen):
    """Aynı anahtara sahip kayıtları sırayla eşleştirir; eşleşmeyenleri döndürür."""
    g1, g2 = defaultdict(list), defaultdict(list)
    for k in l1:
        g1[anahtar_fn(k)].append(k)
    for k in l2:
        g2[anahtar_fn(k)].append(k)
    kalan1, kalan2 = [], []
    for anahtar in list(g1) + [x for x in g2 if x not in g1]:
        x1, x2 = g1.get(anahtar, []), g2.get(anahtar, [])
        for a, b in zip(x1, x2):
            b_maas = b["maas"] + ((b["elden"] or 0) if elden_dahil else 0)
            fark = b_maas - a["maas"]
            eslesen.append({"a": a, "b": b, "b_maas": b_maas, "fark": fark,
                            "ayni": abs(fark) <= tolerans, "yontem": yontem,
                            "kontrol": yontem == "TC" and not isim_uyumlu(a["ad"], b["ad"])})
        kalan1 += x1[len(x2):]
        kalan2 += x2[len(x1):]
    return kalan1, kalan2


def karsilastir(k1, k2, tolerans=0.0, elden_dahil=False):
    """Önce TC kimlik no ile, eşleşmeyenleri isimle eşleştirir.
    Aynı anahtar birden fazlaysa sırayla eşleştirir (1. ile 1., 2. ile 2.)."""
    eslesen = []
    tc1 = [k for k in k1 if k["tc"]]
    tc2 = [k for k in k2 if k["tc"]]
    kalan1, kalan2 = _esle(tc1, tc2, lambda k: k["tc"], "TC", tolerans, elden_dahil, eslesen)
    kalan1 += [k for k in k1 if not k["tc"]]
    kalan2 += [k for k in k2 if not k["tc"]]
    sadece1, sadece2 = _esle(kalan1, kalan2, lambda k: k["anahtar"], "İsim", tolerans, elden_dahil, eslesen)
    return eslesen, sadece1, sadece2


# ---------------------------------------------------------------- Excel çıktısı
BASLIK = PatternFill("solid", fgColor="1F4E78")
KIRMIZI = PatternFill("solid", fgColor="F8CBAD")
YESIL = PatternFill("solid", fgColor="C6E0B4")
SARI = PatternFill("solid", fgColor="FFE699")


def sayfa_yaz(wb, ad, basliklar, satirlar, renkler=None, para_sutunlari=()):
    ws = wb.create_sheet(ad)
    ws.append(basliklar)
    for h in ws[1]:
        h.font = Font(bold=True, color="FFFFFF")
        h.fill = BASLIK
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for n, s in enumerate(satirlar):
        ws.append(s)
        if renkler and renkler[n]:
            for c in ws[ws.max_row]:
                c.fill = renkler[n]
    for col in para_sutunlari:
        for c in ws[col][1:]:
            c.number_format = "#,##0.00"
    for i, h in enumerate(basliklar, 1):
        uz = max([len(str(h))] + [len(str(s[i - 1])) for s in satirlar if i - 1 < len(s)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(uz + 3, 12), 45)
    ws.freeze_panes = "A2"
    if satirlar:
        ws.auto_filter.ref = ws.dimensions
    return ws


def rapor_yaz(yol, ad1, ad2, eslesen, sadece1, sadece2, tolerans, elden_dahil=False):
    wb = Workbook()
    ozet = wb.active
    ozet.title = "Özet"
    farkli = [e for e in eslesen if not e["ayni"]]
    ayni = len(eslesen) - len(farkli)
    ozet_satirlar = [
        ("Dosya 1 (sayfalı)", os.path.basename(ad1)),
        ("Dosya 2 (tek liste)", os.path.basename(ad2)),
        ("Tolerans (TL)", tolerans),
        ("Dosya 2 maaşı", "MAAŞ + ELDEN" if elden_dahil else "Sadece MAAŞ sütunu"),
        ("", ""),
        ("Eşleşen kişi sayısı", len(eslesen)),
        ("  Maaşı aynı olan", ayni),
        ("  Maaşı farklı olan", len(farkli)),
        ("  TC aynı ama İSİM FARKLI (kontrol edin)", sum(e["kontrol"] for e in eslesen)),
        ("  (TC ile eşleşen / isimle eşleşen)",
         f"{sum(e['yontem'] == 'TC' for e in eslesen)} / {sum(e['yontem'] == 'İsim' for e in eslesen)}"),
        ("Sadece Dosya 1'de olan", len(sadece1)),
        ("Sadece Dosya 2'de olan", len(sadece2)),
        ("", ""),
        ("Farkların toplamı (Dosya2 - Dosya1)", round(sum(e["fark"] for e in farkli), 2)),
        ("", ""),
        ("Dosya 1 genel toplam maaş", round(sum(e["a"]["maas"] for e in eslesen) + sum(k["maas"] for k in sadece1), 2)),
        ("Dosya 2 genel toplam maaş",
         round(sum(e["b_maas"] for e in eslesen)
               + sum(k["maas"] + ((k["elden"] or 0) if elden_dahil else 0) for k in sadece2), 2)),
        ("Eşleşenlerin toplamı - Dosya 1", round(sum(e["a"]["maas"] for e in eslesen), 2)),
        ("Eşleşenlerin toplamı - Dosya 2", round(sum(e["b_maas"] for e in eslesen), 2)),
        ("", ""),
        (f"'Asgari' yazılı maaş sayısı (1 asgari = {ASGARI_MAAS:,.2f} TL)",
         sum(e["a"]["asgari"] for e in eslesen) + sum(k["asgari"] for k in sadece1)
         + sum(e["b"]["asgari"] for e in eslesen) + sum(k["asgari"] for k in sadece2)),
    ]
    for s in ozet_satirlar:
        ozet.append(s)
    for r in ozet["A"]:
        r.font = Font(bold=True)
    ozet.column_dimensions["A"].width = 40
    ozet.column_dimensions["B"].width = 40

    tum = sorted(eslesen, key=lambda e: (e["ayni"], -abs(e["fark"])))
    basliklar = ["Okul", "Bölüm / Sayfa", "TC No (Dosya 1)", "TC No (Dosya 2)", "İsim (Dosya 1)", "Maaş (Dosya 1)",
                 "İsim (Dosya 2)", "Maaş (Dosya 2)", "Elden (Dosya 2)", "Karşılaştırılan D2 Maaşı",
                 "Fark (D2 - D1)", "Durum", "Eşleşme", "Kontrol"]

    def satir(e):
        a, b = e["a"], e["b"]
        return [a["okul"], a["sayfa"], a["tc"], b["tc"], a["ad"], a["maas"], b["ad"], b["maas"], b["elden"],
                e["b_maas"], round(e["fark"], 2),
                "Aynı" if e["ayni"] else ("Dosya 2 yüksek" if e["fark"] > 0 else "Dosya 1 yüksek"), e["yontem"],
                "; ".join(x for x in ("İSİM FARKLI - kontrol edin" if e["kontrol"] else "",
                                     "Dosya 1: ASGARİ" if a["asgari"] else "",
                                     "Dosya 2: ASGARİ" if b["asgari"] else "") if x)]

    para = ("F", "H", "I", "J", "K")
    sayfa_yaz(wb, "Tüm Karşılaştırma", basliklar, [satir(e) for e in tum],
              [YESIL if e["ayni"] else KIRMIZI for e in tum], para)
    f = sorted(farkli, key=lambda e: -abs(e["fark"]))
    sayfa_yaz(wb, "Farklı Maaşlar", basliklar, [satir(e) for e in f], [KIRMIZI] * len(f), para)
    kont = [e for e in tum if e["kontrol"]]
    sayfa_yaz(wb, "İsim Kontrol", basliklar, [satir(e) for e in kont], [SARI] * len(kont), para)
    sayfa_yaz(wb, "Sadece Dosya 1'de",
              ["Okul", "Bölüm / Sayfa", "TC No", "İsim", "Maaş", "Satır"],
              [[k["okul"], k["sayfa"], k["tc"], k["ad"], k["maas"], k["satir"]] for k in sadece1],
              [SARI] * len(sadece1), ("E",))
    sayfa_yaz(wb, "Sadece Dosya 2'de",
              ["Sayfa", "TC No", "İsim", "Maaş", "Elden", "Satır"],
              [[k["sayfa"], k["tc"], k["ad"], k["maas"], k["elden"], k["satir"]] for k in sadece2],
              [SARI] * len(sadece2), ("D", "E"))
    wb.save(yol)


# ---------------------------------------------------------------- dosya seçimi
def dosya_sor(baslik, varsayilan=None):
    try:
        import tkinter as tk
        from tkinter import filedialog
        kok = tk.Tk()
        kok.withdraw()
        kok.attributes("-topmost", True)
        yol = filedialog.askopenfilename(title=baslik, filetypes=[("Excel", "*.xlsx *.xlsm")])
        kok.destroy()
        return yol
    except Exception:
        return input(f"{baslik} (dosya yolu): ").strip().strip('"')


def kaydet_sor(varsayilan):
    try:
        import tkinter as tk
        from tkinter import filedialog
        kok = tk.Tk()
        kok.withdraw()
        kok.attributes("-topmost", True)
        yol = filedialog.asksaveasfilename(title="Sonuç dosyasını kaydet", defaultextension=".xlsx",
                                           initialfile=varsayilan, filetypes=[("Excel", "*.xlsx")])
        kok.destroy()
        return yol or varsayilan
    except Exception:
        return varsayilan


def main():
    global ASGARI_MAAS
    p = argparse.ArgumentParser(description="İki Excel dosyasındaki maaşları isme göre karşılaştırır.")
    p.add_argument("dosya1", nargs="?", help="Sayfalara bölünmüş Excel (anaokulu, lise ...)")
    p.add_argument("dosya2", nargs="?", help="Tek listeli Excel")
    p.add_argument("-o", "--cikti", help="Sonuç dosyası (varsayılan: maas_karsilastirma.xlsx)")
    p.add_argument("--tolerans", type=float, default=0.0, help="Bu tutara kadar farkı 'aynı' say (varsayılan 0)")
    p.add_argument("--elden-dahil", action="store_true",
                   help="Dosya 2'de MAAŞ + ELDEN toplamını Dosya 1 maaşıyla karşılaştır")
    p.add_argument("--asgari", type=float, default=ASGARI_MAAS,
                   help="Maaş hücresinde 'asgari' yazan satırlar için tutar (varsayılan 28075)")
    p.add_argument("--ad1", help="Dosya 1 isim sütunu harfi (örn. B) - otomatik bulma yerine")
    p.add_argument("--maas1", help="Dosya 1 maaş sütunu harfi (örn. F)")
    p.add_argument("--ad2", help="Dosya 2 isim sütunu harfi")
    p.add_argument("--maas2", help="Dosya 2 maaş sütunu harfi")
    a = p.parse_args()

    ASGARI_MAAS = a.asgari
    d1 = a.dosya1 or dosya_sor("1) Sayfalara bölünmüş Excel'i seçin")
    d2 = a.dosya2 or dosya_sor("2) Tek listeli Excel'i seçin")
    if not d1 or not d2:
        sys.exit("Dosya seçilmedi.")
    for d in (d1, d2):
        if not os.path.isfile(d):
            sys.exit(f"Dosya bulunamadı: {d}")
        if d.lower().endswith(".xls"):
            sys.exit(f"'{d}' eski .xls biçiminde. Excel'de 'Farklı Kaydet' ile .xlsx yapıp tekrar deneyin.")

    k1, u1 = kisileri_oku(d1, a.ad1, a.maas1, "Dosya 1")
    k2, u2 = kisileri_oku(d2, a.ad2, a.maas2, "Dosya 2")
    for u in u1 + u2:
        print("UYARI:", u)
    if not k1 or not k2:
        sys.exit("Dosyalardan kayıt okunamadı. Başlıkları kontrol edin veya --ad1/--maas1 ... ile sütun belirtin.")

    for etiket, k in (("Dosya 1", k1), ("Dosya 2", k2)):
        sayac = defaultdict(int)
        for x in k:
            sayac[x["anahtar"]] += 1
        tekrar = [x for x, n in sayac.items() if n > 1]
        if tekrar:
            print(f"UYARI: {etiket}'de aynı isimle birden fazla kayıt var ({len(tekrar)} isim); sırayla eşleştirildi.")

    eslesen, s1, s2 = karsilastir(k1, k2, a.tolerans, a.elden_dahil)
    cikti = a.cikti or (kaydet_sor("maas_karsilastirma.xlsx") if not (a.dosya1 and a.dosya2) else "maas_karsilastirma.xlsx")
    rapor_yaz(cikti, d1, d2, eslesen, s1, s2, a.tolerans, a.elden_dahil)

    farkli = sum(1 for e in eslesen if not e["ayni"])
    print(f"Dosya 1: {len(k1)} kişi ({len({k['sayfa'] for k in k1})} sayfa) | Dosya 2: {len(k2)} kişi")
    print(f"Eşleşen: {len(eslesen)} | Maaşı farklı: {farkli} | Sadece D1: {len(s1)} | Sadece D2: {len(s2)}")
    kontrol = sum(e["kontrol"] for e in eslesen)
    if kontrol:
        print(f"DİKKAT: {kontrol} kişide TC aynı ama isim farklı! 'İsim Kontrol' sayfasına bakın.")
    print(f"Sonuç kaydedildi: {os.path.abspath(cikti)}")
    if sys.platform.startswith("win") and not (a.dosya1 and a.dosya2):
        input("Kapatmak için Enter'a basın...")


if __name__ == "__main__":
    main()
