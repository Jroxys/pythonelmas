"""Firma ödeme takip tablosu (ODEME_TAKIBI.xlsx) oluşturur.

Çalıştırmak için:  python odeme_takibi.py
"""
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

DOSYA = "ODEME_TAKIBI.xlsx"
SATIR_SAYISI = 500  # kaç satırlık kayıt alanı hazırlanacak
SON = SATIR_SAYISI + 1

ODENDI = "Ödendi"
ODENMEDI = "Ödenmedi"
BEKLENIYOR = "Ödeme Bekleniyor"

BASLIKLAR = [
    ("FATURA TARİHİ", 18),
    ("FİRMA ADI", 30),
    ("FATURA NO", 16),
    ("TUTAR (TL)", 16),
    ("ÖDENEN (TL)", 16),
    ("ÖDEME DURUMU", 20),
    ("ÖDEME TARİHİ", 18),
    ("KALAN (TL)", 16),
]

TL = '#,##0.00 "₺"'
TARIH = "DD.MM.YYYY"
FONT = "Arial"

ince = Side(style="thin", color="BFBFBF")
kenar = Border(left=ince, right=ince, top=ince, bottom=ince)
baslik_dolgu = PatternFill("solid", start_color="1F4E78")
baslik_font = Font(name=FONT, bold=True, color="FFFFFF")
normal_font = Font(name=FONT)
ortala = Alignment(horizontal="center", vertical="center")

wb = Workbook()
ws = wb.active
ws.title = "Ödeme Takibi"

# Başlıklar
for i, (ad, genislik) in enumerate(BASLIKLAR, start=1):
    h = ws.cell(row=1, column=i, value=ad)
    h.font, h.fill, h.alignment, h.border = baslik_font, baslik_dolgu, ortala, kenar
    ws.column_dimensions[h.column_letter].width = genislik
ws.row_dimensions[1].height = 24

# Örnek kayıtlar (renkleri göstermek için; silinebilir)
ornekler = [
    (date(2026, 9, 1), "Örnek Firma A.Ş.", "FTR-2026-001", 15000, 15000, ODENDI, date(2026, 9, 15)),
    (date(2026, 9, 10), "Örnek Ticaret Ltd.", "FTR-2026-002", 8500, 3000, BEKLENIYOR, date(2026, 10, 20)),
    (date(2026, 8, 5), "Örnek İnşaat", "FTR-2026-003", 22750.5, 0, ODENMEDI, date(2026, 9, 5)),
]
for r, kayit in enumerate(ornekler, start=2):
    for c, deger in enumerate(kayit, start=1):
        ws.cell(row=r, column=c, value=deger)

# Biçimler ve formüller
for r in range(2, SON + 1):
    for c in range(1, 9):
        cell = ws.cell(row=r, column=c)
        cell.font, cell.border = normal_font, kenar
    ws.cell(row=r, column=1).number_format = TARIH
    ws.cell(row=r, column=7).number_format = TARIH
    for c in (1, 3, 6, 7):
        ws.cell(row=r, column=c).alignment = ortala
    for c in (4, 5, 8):
        ws.cell(row=r, column=c).number_format = TL
    # KALAN = TUTAR - ÖDENEN (tutar boşsa boş kalır)
    ws.cell(row=r, column=8, value=f'=IF(D{r}="","",D{r}-E{r})')

# Tarih girişleri (A ve G sütunları)
tarih_dv = DataValidation(
    type="date", operator="greaterThan", formula1="DATE(2000,1,1)", allow_blank=True,
    promptTitle="Tarih", prompt="Tarihi GG.AA.YYYY şeklinde girin (örn. 15.10.2026).\nBugünün tarihi: Ctrl + ;",
    errorTitle="Geçersiz tarih", error="Lütfen geçerli bir tarih girin (GG.AA.YYYY).",
)
tarih_dv.add(f"A2:A{SON}")
tarih_dv.add(f"G2:G{SON}")

# Tutar girişleri (D ve E sütunları)
tutar_dv = DataValidation(
    type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
    errorTitle="Geçersiz tutar", error="Tutar 0 veya daha büyük bir sayı olmalı.",
)
tutar_dv.add(f"D2:E{SON}")

# Ödeme durumu açılır listesi (F sütunu)
durum_dv = DataValidation(
    type="list", formula1=f'"{ODENDI},{ODENMEDI},{BEKLENIYOR}"', allow_blank=True,
    promptTitle="Ödeme Durumu", prompt="Listeden seçin.",
    errorTitle="Geçersiz seçim", error="Lütfen listeden bir durum seçin.",
)
durum_dv.add(f"F2:F{SON}")
for dv in (tarih_dv, tutar_dv, durum_dv):
    ws.add_data_validation(dv)

# Satır renklendirme: yeşil = ödendi, kırmızı = ödenmedi, sarı = bekleniyor
renkler = [
    (ODENDI, "C6EFCE", "006100"),
    (ODENMEDI, "FFC7CE", "9C0006"),
    (BEKLENIYOR, "FFEB9C", "9C5700"),
]
for durum, dolgu, yazi in renkler:
    ws.conditional_formatting.add(
        f"A2:H{SON}",
        FormulaRule(
            formula=[f'$F2="{durum}"'],
            fill=PatternFill("solid", start_color=dolgu, end_color=dolgu),
            font=Font(color=yazi),
        ),
    )

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:H{SON}"

# Özet tablosu (J-K sütunları)
ws.column_dimensions["I"].width = 3
ws.column_dimensions["J"].width = 26
ws.column_dimensions["K"].width = 18
ozet_baslik = ws.cell(row=1, column=10, value="ÖZET")
ws.merge_cells("J1:K1")
ozet_baslik.font, ozet_baslik.fill, ozet_baslik.alignment = baslik_font, baslik_dolgu, ortala

aralik = lambda s: f"{s}2:{s}{SON}"
ozet = [
    ("Toplam Fatura Tutarı", f"=SUM({aralik('D')})", TL, None),
    ("Toplam Ödenen", f"=SUM({aralik('E')})", TL, None),
    ("Toplam Kalan", f"=SUM({aralik('H')})", TL, None),
    ("Ödendi (adet)", f'=COUNTIF({aralik("F")},"{ODENDI}")', "0", "C6EFCE"),
    ("Ödenmedi (adet)", f'=COUNTIF({aralik("F")},"{ODENMEDI}")', "0", "FFC7CE"),
    ("Ödeme Bekleniyor (adet)", f'=COUNTIF({aralik("F")},"{BEKLENIYOR}")', "0", "FFEB9C"),
    ("Ödenmeyenlerin Kalanı", f'=SUMIF({aralik("F")},"{ODENMEDI}",{aralik("H")})', TL, "FFC7CE"),
    ("Beklenenlerin Kalanı", f'=SUMIF({aralik("F")},"{BEKLENIYOR}",{aralik("H")})', TL, "FFEB9C"),
]
for r, (etiket, formul, bicim, renk) in enumerate(ozet, start=2):
    e = ws.cell(row=r, column=10, value=etiket)
    v = ws.cell(row=r, column=11, value=formul)
    e.font, v.font = Font(name=FONT, bold=True), normal_font
    e.border = v.border = kenar
    v.number_format = bicim
    if renk:
        e.fill = PatternFill("solid", start_color=renk)

# Kullanım notu
not_satiri = len(ozet) + 3
ws.cell(row=not_satiri, column=10, value="NASIL KULLANILIR?").font = Font(name=FONT, bold=True)
notlar = [
    "• A–G sütunlarına bilgileri girin.",
    "• KALAN (TL) otomatik hesaplanır: Tutar − Ödenen.",
    "• Ödeme Durumu listeden seçilir; satır rengi",
    "  otomatik değişir (yeşil / kırmızı / sarı).",
    "• Tarihler GG.AA.YYYY şeklinde; bugün için Ctrl + ;",
    "• 2–4. satırlar örnektir, silebilirsiniz.",
]
for i, metin in enumerate(notlar, start=1):
    ws.cell(row=not_satiri + i, column=10, value=metin).font = Font(name=FONT, size=9, italic=True)

ws["H1"].comment = Comment("Otomatik hesaplanır: TUTAR − ÖDENEN. Elle yazmayın.", "Ödeme Takibi")

wb.save(DOSYA)
print(f"{DOSYA} oluşturuldu.")
